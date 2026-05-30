"""Generate a PowerPoint deck (paper/slides.pptx) for the G2P PEFT study.

Renders the figures to PNG (PowerPoint cannot embed PDF), builds a native
results table from results/runs, and assembles a 16:9 deck mirroring the paper.

    python scripts/make_pptx.py
"""
from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

from analysis.schema import load_results
from analysis.aggregate import aggregate, cells_by_key
from analysis.constants import METHOD_ORDER, DISPLAY_METHOD, LANG_ORDER, SIGMORPHON_BASELINE_WER
from analysis.figures import make_pareto_figure, make_forgetting_figure

ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / "figures"
REPO = "https://github.com/stoianraresstefan/G2P-Parameter-Efficient-Fine-Tuning"

TITLE_C = RGBColor(0x1F, 0x3A, 0x6E)   # dark blue
ACCENT_C = RGBColor(0xB0, 0x30, 0x2A)  # deep red
DARK = RGBColor(0x22, 0x22, 0x22)
HEADER_BG = RGBColor(0x1F, 0x3A, 0x6E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xEC, 0xEF, 0xF5)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = prs.slide_width, prs.slide_height


def slide():
    return prs.slides.add_slide(BLANK)


def _runs_from_markup(paragraph, text, size, color=DARK):
    """Split on ** to alternate normal / bold runs."""
    parts = text.split("**")
    for i, part in enumerate(parts):
        if part == "":
            continue
        r = paragraph.add_run()
        r.text = part
        r.font.size = Pt(size)
        r.font.bold = i % 2 == 1
        r.font.color.rgb = ACCENT_C if i % 2 == 1 else color


def add_title(s, text, size=30):
    tb = s.shapes.add_textbox(Inches(0.6), Inches(0.3), Inches(12.1), Inches(1.0))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.color.rgb = TITLE_C
    # accent underline bar
    bar = s.shapes.add_shape(1, Inches(0.65), Inches(1.25), Inches(2.2), Pt(3))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT_C
    bar.line.fill.background()
    return s


def add_bullets(s, items, top=1.55, left=0.8, width=11.8, height=5.4, size=20):
    tb = s.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    first = True
    for item in items:
        if isinstance(item, tuple):
            level, text = item
        else:
            level, text = 0, item
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(10)
        p.space_before = Pt(2)
        prefix = "•  " if level == 0 else "      –  "
        lead = p.add_run()
        lead.text = prefix
        lead.font.size = Pt(size if level == 0 else size - 2)
        lead.font.color.rgb = ACCENT_C if level == 0 else DARK
        _runs_from_markup(p, text, size if level == 0 else size - 2)
    return s


# ---------------------------------------------------------------- title slide
s = slide()
# accent band
band = s.shapes.add_shape(1, 0, Inches(2.4), SW, Inches(0.06))
band.fill.solid(); band.fill.fore_color.rgb = ACCENT_C; band.line.fill.background()
tb = s.shapes.add_textbox(Inches(0.8), Inches(0.9), Inches(11.7), Inches(1.5))
tf = tb.text_frame; tf.word_wrap = True
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "LoRA for Multilingual Grapheme-to-Phoneme Conversion"
r.font.size = Pt(34); r.font.bold = True; r.font.color.rgb = TITLE_C
tb2 = s.shapes.add_textbox(Inches(0.8), Inches(2.6), Inches(11.7), Inches(0.9))
tf2 = tb2.text_frame; tf2.word_wrap = True
p = tf2.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "Parameter-Efficient Adaptation of ByT5-G2P on Romanian and Modern Greek"
r.font.size = Pt(19); r.font.italic = True; r.font.color.rgb = DARK
tb3 = s.shapes.add_textbox(Inches(0.8), Inches(4.0), Inches(11.7), Inches(2.0))
tf3 = tb3.text_frame; tf3.word_wrap = True
for txt, sz, bold in [
    ("Dan-Gabriel Tudor   ·   Ştefan-Rareş Stoian", 22, True),
    ("Faculty of Mathematics and Informatics, University of Bucharest", 16, False),
    ("LLM for NLP — 2026", 14, False),
    (REPO, 12, False),
]:
    p = tf3.add_paragraph(); p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = txt; r.font.size = Pt(sz); r.font.bold = bold
    r.font.color.rgb = TITLE_C if bold else DARK

# ---------------------------------------------------------------- G2P task
s = slide(); add_title(s, "The grapheme-to-phoneme (G2P) task")
add_bullets(s, [
    "**G2P** maps spelling to pronunciation — a sequence of phonemes (IPA).",
    "It is the front-end of **text-to-speech** and supplies the lexicons behind **speech recognition**.",
    "Example: Romanian *porneşte* (“[it] starts”) →  p o r n e ʃ t e,  where ⟨ş⟩ → /ʃ/.",
    "Spelling-to-sound mappings are **irregular** even in shallow orthographies — so G2P is nontrivial.",
])

# ---------------------------------------------------------------- why multilingual
s = slide(); add_title(s, "Why multilingual G2P matters")
add_bullets(s, [
    "High-quality pronunciation lexicons need linguistic expertise and are **scarce** for most languages.",
    "Pretrained multilingual models (CharsiuG2P / ByT5-G2P) transfer across **~100 languages** from a **single** network…",
    "… and need only an ISO-639-3 language tag to transcribe a word.",
    "Strong off the shelf — but still benefit from **per-language adaptation** when in-language data exists.",
])

# ---------------------------------------------------------------- cost problem
s = slide(); add_title(s, "The cost problem")
add_bullets(s, [
    "Adapting by **full fine-tuning** stores a full copy of ~300M weights **per language**:  K languages ⇒ K × 300M.",
    "It can also **overwrite** previously learned languages (catastrophic forgetting).",
    "**Parameter-efficient fine-tuning (PEFT)** — LoRA, adapters — is standard for LLMs and ASR…",
    "… but has **not** been adopted for G2P.   **This is the gap we study.**",
])

# ---------------------------------------------------------------- research question
s = slide(); add_title(s, "Research question")
box = s.shapes.add_shape(1, Inches(1.2), Inches(2.2), Inches(10.9), Inches(2.2))
box.fill.solid(); box.fill.fore_color.rgb = LIGHT; box.line.color.rgb = TITLE_C; box.line.width = Pt(1.5)
tf = box.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
tf.margin_left = Inches(0.4); tf.margin_right = Inches(0.4)
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
_runs_from_markup(p, "Can **LoRA** match **full fine-tuning** of ByT5-G2P on low-resource "
                     "languages while training **under 1%** of the parameters — and how do "
                     "they compare on **catastrophic forgetting**?", 24)
add_bullets(s, [
    "First systematic study of LoRA for multilingual G2P.",
    "Evaluated on Romanian and Modern Greek (SIGMORPHON 2021, low-resource).",
], top=4.7, size=18)

# ---------------------------------------------------------------- method
s = slide(); add_title(s, "Method: one model, four adaptation strategies")
add_bullets(s, [
    "Base model: **CharsiuG2P ByT5-small** (~300M params, byte-level, no language-specific vocabulary).",
    (1, "**Zero-shot** — pretrained model, no training (0%)."),
    (1, "**Frozen head** — train only the LM head (0.19%)."),
    (1, "**LoRA** — low-rank updates on attention q/k/v/o, r ∈ {8,16}, α = 2r  (0.39–0.78%):  W + (α/r)·BA."),
    (1, "**Full fine-tuning** — all parameters (100%)."),
    "Greedy decoding; input prompt  <ron>:  /  <gre>: .",
], size=19)

# ---------------------------------------------------------------- data & setup
s = slide(); add_title(s, "Data and experimental setup")
add_bullets(s, [
    "**SIGMORPHON 2021**, low-resource: Romanian (Latin / Romance) and Greek (Greek script / Hellenic).",
    (1, "800 / 100 / 100 words per language; IPA targets."),
    "**Metrics:** PER (character-level edit rate) and WER; **forgetting** on held-out English (eng-us) IPA.",
    "**Compute:** Kaggle dual T4 GPUs; HuggingFace transformers + peft.",
    "Two fixes that mattered: the tag **<ron>** (not file code rum: 5.96 vs 24.76 dev PER), and **character-level** scoring so the un-delimited zero-shot output is comparable.",
], size=19)

# ---------------------------------------------------------------- main results table
records = load_results(ROOT / "results" / "runs")
cells = aggregate(records)
by = cells_by_key(cells)

# render figures to PNG for the picture slides
make_pareto_figure(cells, FIG / "fig1_pareto.png")
make_forgetting_figure(cells, FIG / "fig2_forgetting.png")

# best (min mean) per (lang, metric)
best = {}
for lang in LANG_ORDER:
    for metric in ("per", "wer"):
        vals = [(c, getattr(by[(lang, c)], f"{metric}_mean")) for c in METHOD_ORDER if (lang, c) in by]
        best[(lang, metric)] = min(vals, key=lambda t: t[1])[0]


def fmt(cell, metric):
    mean = getattr(cell, f"{metric}_mean")
    std = getattr(cell, f"{metric}_std")
    if cell.n_seeds > 1 and std > 0:
        return f"{mean:.2f} ± {std:.2f}"
    return f"{mean:.2f}"


s = slide(); add_title(s, "Main results")
n_rows = 1 + len(METHOD_ORDER) + 1
gt = s.shapes.add_table(n_rows, 5, Inches(1.1), Inches(1.5), Inches(11.1), Inches(3.4)).table
headers = ["Method", "Ro PER", "Ro WER", "El PER", "El WER"]
gt.columns[0].width = Inches(3.5)
for ci in range(1, 5):
    gt.columns[ci].width = Inches(1.9)
for ci, h in enumerate(headers):
    c = gt.cell(0, ci); c.text = h
    pr = c.text_frame.paragraphs[0]; pr.runs[0].font.bold = True
    pr.runs[0].font.size = Pt(15); pr.runs[0].font.color.rgb = WHITE
    pr.alignment = PP_ALIGN.CENTER if ci else PP_ALIGN.LEFT
    c.fill.solid(); c.fill.fore_color.rgb = HEADER_BG
for ri, cfg in enumerate(METHOD_ORDER, start=1):
    vals = [DISPLAY_METHOD[cfg]]
    flags = [False]
    for lang in LANG_ORDER:
        for metric in ("per", "wer"):
            cell = by.get((lang, cfg))
            vals.append(fmt(cell, metric) if cell else "--")
            flags.append(best.get((lang, metric)) == cfg)
    for ci, (v, bold) in enumerate(zip(vals, flags)):
        c = gt.cell(ri, ci); c.text = v
        pr = c.text_frame.paragraphs[0]
        pr.runs[0].font.size = Pt(14); pr.runs[0].font.bold = bold
        pr.runs[0].font.color.rgb = ACCENT_C if bold else DARK
        pr.alignment = PP_ALIGN.CENTER if ci else PP_ALIGN.LEFT
        c.fill.solid(); c.fill.fore_color.rgb = LIGHT if ri % 2 else WHITE
# baseline row
br = n_rows - 1
base_vals = ["SIGMORPHON'21 WER", "--", f"{SIGMORPHON_BASELINE_WER['rum']:.2f}", "--", f"{SIGMORPHON_BASELINE_WER['gre']:.2f}"]
for ci, v in enumerate(base_vals):
    c = gt.cell(br, ci); c.text = v
    pr = c.text_frame.paragraphs[0]; pr.runs[0].font.size = Pt(13); pr.runs[0].font.italic = True
    pr.runs[0].font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    pr.alignment = PP_ALIGN.CENTER if ci else PP_ALIGN.LEFT
    c.fill.solid(); c.fill.fore_color.rgb = WHITE
add_bullets(s, [
    "**Romanian:** adaptation helps — full FT best (PER 1.76, WER 10); LoRA r=16 (3.36) beats zero-shot (4.48) at **0.78%** of params.",
    "**Greek:** zero-shot already excellent (PER 1.56); no fine-tuning improves it.",
], top=5.25, size=16)

# ---------------------------------------------------------------- pareto figure
s = slide(); add_title(s, "Accuracy vs. cost: the Pareto view")
s.shapes.add_picture(str(FIG / "fig1_pareto.png"), Inches(1.4), Inches(1.6), width=Inches(10.5))
tb = s.shapes.add_textbox(Inches(0.8), Inches(6.55), Inches(11.8), Inches(0.7))
p = tb.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
_runs_from_markup(p, "Romanian: each extra parameter budget buys lower PER.    "
                     "Greek: **zero-shot dominates** the frontier.", 15)

# ---------------------------------------------------------------- forgetting figure
s = slide(); add_title(s, "Catastrophic forgetting on English")
s.shapes.add_picture(str(FIG / "fig2_forgetting.png"), Inches(0.7), Inches(1.7), height=Inches(5.0))
add_bullets(s, [
    "Base model: PER **36.3**.",
    "**Full fine-tuning forgets most** (Greek-adapted 43.3, +7).",
    "**LoRA / frozen-head stay near baseline.**",
    "Frozen base ⇒ multilingual ability preserved.",
    "The cleanest argument for **PEFT**.",
], top=2.0, left=6.6, width=6.2, size=18)

# ---------------------------------------------------------------- limitations
s = slide(); add_title(s, "Limitations and future work")
add_bullets(s, [
    "**Scope:** two languages, two LoRA ranks, mostly single-seed, one base model, isolated words, no adapter baseline.",
    "**Threat to validity:** Greek zero-shot likely inflated by WikiPron/SIGMORPHON pretraining overlap (within-split leakage is 0% exact).",
    "**Metric:** character-level PER is internally consistent but not the official phoneme-token PER.",
    "**Future:** adapters and wider rank sweep; more languages/scripts; self-training; sentence-level G2P; robustness.",
], size=19)

# ---------------------------------------------------------------- conclusion
s = slide(); add_title(s, "Conclusion")
add_bullets(s, [
    "LoRA **approaches full fine-tuning** where adaptation helps (Romanian PER 4.48 → 3.36) at **<1%** of the parameters.",
    "Per-language storage drops from K × ~300M to K × ~1–2M weights.",
    "By freezing the backbone, LoRA **preserves the multilingual base** where full fine-tuning forgets it.",
], top=1.6, size=20)
box = s.shapes.add_shape(1, Inches(1.2), Inches(4.7), Inches(10.9), Inches(1.5))
box.fill.solid(); box.fill.fore_color.rgb = TITLE_C; box.line.fill.background()
tf = box.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "Takeaway: PEFT is a practical default for per-language G2P deployment."
r.font.size = Pt(22); r.font.bold = True; r.font.color.rgb = WHITE

# ---------------------------------------------------------------- thank you
s = slide()
tb = s.shapes.add_textbox(Inches(1.0), Inches(2.9), Inches(11.3), Inches(1.8))
tf = tb.text_frame; tf.word_wrap = True
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "Thank you — Questions?"; r.font.size = Pt(40); r.font.bold = True; r.font.color.rgb = TITLE_C
p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
r = p2.add_run(); r.text = REPO; r.font.size = Pt(14); r.font.color.rgb = DARK

out = ROOT / "paper" / "slides.pptx"
prs.save(str(out))
print(f"Wrote {out}  ({sum(1 for _ in prs.slides)} slides)")
